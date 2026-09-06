from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable


DIRECT = "可直接解析"
SPECIAL = "需特殊处理"
UNAVAILABLE = "暂无法解析"


@dataclass(frozen=True)
class ParseResult:
    status: str
    content_length: int = 0
    document_scale: int = 0
    scale_unit: str = ""
    issue: str = ""
    text: str = ""
    layout_warnings: tuple[str, ...] = ()


def normalize_text(parts: list[str]) -> str:
    return "\n".join(part.strip() for part in parts if part and part.strip())


def count_content_characters(text: str) -> int:
    """统计用户易理解的正文字符数，忽略空白和纯标点。"""
    return len(re.findall(r"[\u3400-\u9fffA-Za-z0-9]", text))


def parse_docx(path: Path) -> ParseResult:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = normalize_text(parts)
    return ParseResult(DIRECT, count_content_characters(text), text=text)


def parse_pdf(path: Path) -> ParseResult:
    from pypdf import PdfReader

    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception:
            unlocked = 0
        if not unlocked:
            return ParseResult(SPECIAL, issue="文件已加密，需要提供可读取版本")

    page_count = len(reader.pages)
    parts: list[str] = []
    failed_pages = 0
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            failed_pages += 1
    text = normalize_text(parts)
    length = count_content_characters(text)

    if page_count and length < max(20, page_count * 10):
        return ParseResult(
            SPECIAL,
            length,
            page_count,
            "页",
            "疑似扫描文件，需要先识别文字",
            text,
        )
    if failed_pages:
        return ParseResult(
            SPECIAL,
            length,
            page_count,
            "页",
            f"有{failed_pages}页文字读取异常，建议重新导出文件",
            text,
        )
    return ParseResult(DIRECT, length, page_count, "页", text=text)


def parse_xlsx(path: Path) -> ParseResult:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    parts: list[str] = []
    formula_count = 0
    preheader_sheet_count = 0
    try:
        for worksheet in workbook.worksheets:
            parts.append(worksheet.title)
            early_row_counts: list[tuple[int, int]] = []
            for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                nonempty_count = 0
                for cell in row:
                    if cell.value is not None:
                        parts.append(str(cell.value))
                        nonempty_count += 1
                    if cell.data_type == "f":
                        formula_count += 1
                if row_number <= 10:
                    early_row_counts.append((row_number, nonempty_count))

            nonempty_rows = [(row, count) for row, count in early_row_counts if count]
            if nonempty_rows:
                first_row, first_count = nonempty_rows[0]
                later_has_table_header = any(count >= 2 for _, count in nonempty_rows[1:4])
                if later_has_table_header and (first_row > 1 or first_count == 1):
                    preheader_sheet_count += 1
        text = normalize_text(parts)
        header_merge_count = 0
        vertical_header_merge = False
        merge_pattern = re.compile(
            rb'<mergeCell\s+ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"'
        )
        with zipfile.ZipFile(path) as archive:
            for member in archive.namelist():
                if not member.startswith("xl/worksheets/sheet") or not member.endswith(".xml"):
                    continue
                for match in merge_pattern.finditer(archive.read(member)):
                    _, start_row_raw, _, end_row_raw = match.groups()
                    start_row, end_row = int(start_row_raw), int(end_row_raw)
                    if start_row <= 10:
                        header_merge_count += 1
                        vertical_header_merge = vertical_header_merge or end_row > start_row

        warnings: list[str] = []
        # 合并标题通常是一行只有一个值，随后才出现多列表头；首行仍有多个字段时，
        # 即使只有一处横向合并，也更像分组表头，应按复杂表格处理。
        if vertical_header_merge or header_merge_count > preheader_sheet_count:
            warnings.append("前10行存在多处或跨行合并单元格，可能是多行表头")
        if preheader_sheet_count:
            warnings.append("正式表头前存在合并标题、说明或空行")
        if formula_count:
            warnings.append(f"包含{formula_count}个计算公式")
        return ParseResult(
            DIRECT,
            count_content_characters(text),
            len(workbook.worksheets),
            "个工作表",
            text=text,
            layout_warnings=tuple(warnings),
        )
    finally:
        workbook.close()


def parse_pptx(path: Path) -> ParseResult:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    parts: list[str] = []
    structured_slides: set[int] = set()
    multi_text_slides: set[int] = set()
    mixed_slides: set[int] = set()
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_shape_count = 0
        picture_count = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
                if shape.text.strip():
                    text_shape_count += 1
            if getattr(shape, "has_table", False):
                structured_slides.add(slide_number)
                for row in shape.table.rows:
                    parts.extend(cell.text for cell in row.cells)
            if getattr(shape, "has_chart", False) or shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                structured_slides.add(slide_number)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
        if text_shape_count >= 5:
            multi_text_slides.add(slide_number)
        if picture_count >= 2 and text_shape_count >= 2:
            mixed_slides.add(slide_number)
    text = normalize_text(parts)
    warnings: list[str] = []
    if structured_slides:
        warnings.append(f"{len(structured_slides)}页包含表格、图表或组合对象")
    if multi_text_slides:
        warnings.append(f"{len(multi_text_slides)}页包含较多独立文本框")
    if mixed_slides:
        warnings.append(f"{len(mixed_slides)}页存在复杂图文混排")
    return ParseResult(
        DIRECT,
        count_content_characters(text),
        len(presentation.slides),
        "页",
        text=text,
        layout_warnings=tuple(warnings),
    )


def parse_text(path: Path) -> ParseResult:
    raw = path.read_bytes()
    if b"\x00" in raw[:4096] and not (raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff")):
        return ParseResult(SPECIAL, issue="文字编码异常，建议重新保存文件")

    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            text = raw.decode(encoding)
            return ParseResult(DIRECT, count_content_characters(text), text=text)
        except UnicodeDecodeError:
            continue
    return ParseResult(SPECIAL, issue="文字编码异常，建议重新保存文件")


PARSERS: dict[str, Callable[[Path], ParseResult]] = {
    ".docx": parse_docx,
    ".pdf": parse_pdf,
    ".xlsx": parse_xlsx,
    ".pptx": parse_pptx,
    ".txt": parse_text,
    ".md": parse_text,
}


def parse_document(path: Path) -> ParseResult:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        if path.suffix.lower() in {".doc", ".xls", ".ppt"}:
            return ParseResult(SPECIAL, issue="旧版Office文件，建议转换为常用格式")
        return ParseResult(UNAVAILABLE, issue="第一版暂不支持读取此类文件")

    try:
        return parser(path)
    except ImportError:
        return ParseResult(SPECIAL, issue="缺少本地解析组件，请重新安装完整工具")
    except PermissionError:
        return ParseResult(SPECIAL, issue="文件无法访问，请检查权限或是否被占用")
    except (OSError, ValueError, KeyError, EOFError, zipfile.BadZipFile) as error:
        return ParseResult(SPECIAL, issue=f"文件可能损坏或格式异常（{type(error).__name__}）")
    except Exception as error:  # 单个复杂文档不能中断整批体检
        return ParseResult(SPECIAL, issue=f"文件读取异常（{type(error).__name__}）")
