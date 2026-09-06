from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .scanner import ScanResult


BRAND = "申悦｜悦览杉合科技有限公司"
POSITIONING = "专注企业知识库规划、知识治理与RAG项目落地"
CONTACT = "15811157318"

BLUE = "1746A2"
LIGHT_BLUE = "EAF2FF"
HEADER_BLUE = "1D4E9E"
GREEN = "DCFCE7"
ORANGE = "FEF3C7"
RED = "FEE2E2"
GRAY = "F1F5F9"
TEXT = "1E293B"
MUTED = "64748B"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="DCE5F0")


def safe_text(value: object) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def add_title(ws, title: str, subtitle: str, end_column: int) -> None:
    end_letter = ws.cell(1, end_column).column_letter
    ws.merge_cells(f"A1:{end_letter}1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells(f"A2:{end_letter}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Microsoft YaHei", size=10, color=MUTED)
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 24


def style_header(row) -> None:
    for cell in row:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=HEADER_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)


def style_data_sheet(ws, *, widths: list[int], header_row: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{ws.cell(ws.max_row, ws.max_column).coordinate}"
    style_header(ws[header_row])
    ws.row_dimensions[header_row].height = 30
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_GRAY)
        ws.row_dimensions[row[0].row].height = 36

    # 保留普通自动筛选，不再同时创建 Excel Table。
    # 两种筛选覆盖同一区域时，部分 Windows Excel 会修复并删除 table1.xml。


def add_status_rules(ws, status_column: str, start_row: int, end_row: int) -> None:
    target = f"{status_column}{start_row}:{status_column}{end_row}"
    for status, color in (
        ("可直接使用", GREEN),
        ("整理后使用", ORANGE),
        ("暂不能使用", RED),
    ):
        ws.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'${status_column}{start_row}="{status}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )


def add_priority_rules(ws, priority_column: str, start_row: int, end_row: int) -> None:
    target = f"{priority_column}{start_row}:{priority_column}{end_row}"
    for priority, color in (
        ("优先处理", RED),
        ("建议处理", ORANGE),
        ("后续优化", GRAY),
    ):
        ws.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'${priority_column}{start_row}="{priority}"'],
                fill=PatternFill("solid", fgColor=color),
            ),
        )


def create_rectification_workbook(result: ScanResult, company_name: str) -> Workbook:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "整改概览"
    overview.sheet_view.showGridLines = False
    overview.column_dimensions["A"].width = 20
    overview.column_dimensions["B"].width = 22
    overview.column_dimensions["C"].width = 24
    overview.column_dimensions["D"].width = 30
    overview.column_dimensions["E"].width = 20
    add_title(overview, "文档整改清单", f"{company_name}｜{BRAND}", 5)

    problematic = [record for record in result.records if record.issues]
    priorities = Counter(record.priority for record in problematic)
    overview["A4"] = "需要整改的资料"
    overview["B4"] = len(problematic)
    overview["C4"] = "优先处理"
    overview["D4"] = priorities["优先处理"]
    overview["E4"] = f"共盘点 {result.total_files:,} 份"
    for cell in overview[4]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(name="Microsoft YaHei", size=11, bold=cell.column in (2, 4), color=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    overview.row_dimensions[4].height = 34

    overview["A6"] = "处理顺序"
    overview["B6"] = "数量"
    overview["C6"] = "怎么理解"
    overview.merge_cells("C6:E6")
    style_header(overview[6])
    descriptions = (
        ("优先处理", priorities["优先处理"], "不处理就无法正常使用，建议首先解决。", RED),
        ("建议处理", priorities["建议处理"], "可能影响知识库效果或产生错误口径。", ORANGE),
        ("后续优化", priorities["后续优化"], "不影响第一阶段建设，可以逐步完善。", GRAY),
    )
    for row_index, (label, count, description, color) in enumerate(descriptions, start=7):
        overview.cell(row_index, 1, label)
        overview.cell(row_index, 2, count)
        overview.cell(row_index, 3, description)
        overview.merge_cells(start_row=row_index, start_column=3, end_row=row_index, end_column=5)
        for cell in overview[row_index]:
            cell.font = Font(name="Microsoft YaHei", size=10, color=TEXT)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        overview.row_dimensions[row_index].height = 28

    overview.merge_cells("A12:E12")
    overview["A12"] = "下一步：知识治理"
    overview["A12"].font = Font(name="Microsoft YaHei", size=14, bold=True, color=WHITE)
    overview["A12"].fill = PatternFill("solid", fgColor=BLUE)
    overview["A12"].alignment = Alignment(vertical="center")
    overview.row_dimensions[12].height = 30
    overview.merge_cells("A13:E15")
    overview["A13"] = (
        "知识库建设，不是只把文件导入系统。在这之前，还要完成版本确认、资料清理、格式处理、"
        "使用范围划分和知识库规划。\n\n"
        "免费预约一次30分钟报告解读，帮您确认哪些问题最值得先处理，以及知识库项目应该从哪里开始。"
    )
    overview["A13"].font = Font(name="Microsoft YaHei", size=10, color=TEXT)
    overview["A13"].alignment = Alignment(vertical="top", wrap_text=True)
    overview["A13"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    overview.merge_cells("A17:E18")
    overview["A17"] = (
        f"{BRAND}\n{POSITIONING}\n手机/微信：{CONTACT}｜微信公众号：互联网悦读笔记｜小红书：申老师的AI赋能笔记"
    )
    overview["A17"].font = Font(name="Microsoft YaHei", size=10, bold=True, color=BLUE)
    overview["A17"].alignment = Alignment(vertical="center", wrap_text=True)
    overview.row_dimensions[17].height = 34
    overview.row_dimensions[18].height = 34

    sheet = workbook.create_sheet("整改清单")
    headers = ["文件名称", "所在位置", "当前情况", "发现的问题", "建议怎么处理", "处理顺序", "建议知识库"]
    sheet.append(headers)
    for record in problematic:
        sheet.append(
            [
                safe_text(record.file_name),
                safe_text(record.relative_path),
                record.use_status,
                safe_text(record.issues),
                safe_text(record.recommendation),
                record.priority,
                record.suggested_knowledge_base,
            ]
        )
    style_data_sheet(sheet, widths=[28, 46, 15, 42, 48, 14, 24], header_row=1)
    if sheet.max_row >= 2:
        add_status_rules(sheet, "C", 2, sheet.max_row)
        add_priority_rules(sheet, "F", 2, sheet.max_row)
    return workbook


def create_inventory_workbook(result: ScanResult, company_name: str) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "完整文档台账"
    add_title(sheet, "完整文档台账", f"{company_name}｜共 {result.total_files:,} 份资料｜{BRAND}", 8)
    headers = ["文件名称", "所在位置", "文件类型", "文件大小（MB）", "最后修改时间", "当前情况", "发现的问题", "建议知识库"]
    sheet.append([])
    sheet.append(headers)
    for record in result.records:
        modified = datetime.fromisoformat(record.modified_at).replace(tzinfo=None)
        sheet.append(
            [
                safe_text(record.file_name),
                safe_text(record.relative_path),
                record.display_type,
                record.size_bytes / (1024 * 1024),
                modified,
                record.use_status,
                safe_text(record.issues),
                record.suggested_knowledge_base,
            ]
        )
    style_data_sheet(sheet, widths=[28, 48, 18, 16, 20, 15, 44, 24], header_row=4)
    for cell in sheet[4][3:5]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(5, sheet.max_row + 1):
        sheet.cell(row, 4).number_format = '0.00 "MB"'
        sheet.cell(row, 4).alignment = Alignment(horizontal="right", vertical="top")
        sheet.cell(row, 5).number_format = "yyyy-mm-dd hh:mm"
        sheet.cell(row, 5).alignment = Alignment(horizontal="center", vertical="top")
    if sheet.max_row >= 5:
        add_status_rules(sheet, "F", 5, sheet.max_row)
    return workbook


def export_workbooks(
    result: ScanResult, output_folder: str | Path, company_name: str | None = None
) -> tuple[Path, Path]:
    output = Path(output_folder).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    company = (company_name or result.root.name or "企业").strip()

    rectification_path = output / "文档整改清单.xlsx"
    inventory_path = output / "完整文档台账.xlsx"
    create_rectification_workbook(result, company).save(rectification_path)
    create_inventory_workbook(result, company).save(inventory_path)

    # 保存后立即重开，提前发现不完整或损坏的工作簿。
    for path in (rectification_path, inventory_path):
        checked = load_workbook(path, read_only=True, data_only=False)
        checked.close()
    return rectification_path, inventory_path
